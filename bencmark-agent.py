from __future__ import annotations

import base64
import io
import json
import shutil
import time
import uuid
from pathlib import Path
from typing import Optional

import pyodbc
from langchain_core.messages import HumanMessage, SystemMessage

from agentcore.custom import Node
from agentcore.io import (
    BoolInput,
    DataInput,
    HandleInput,
    IntInput,
    MessageInput,
    MessageTextInput,
    MultilineInput,
    Output,
)
from agentcore.schema.data import Data
from agentcore.schema.message import Message

# ── Prompts ───────────────────────────────────────────────────────────────────

_CLASSIFY_SYSTEM = """You are an expert document classifier for Motherson's procurement team. Classify files into EXACTLY ONE of six categories.

CATEGORIES & MANDATORY FIELDS:
1. MPBC — Motherson Purchase BID Comparison: 3+ suppliers compared side-by-side, RAS Number, GSP terminology, MPBC/Preisspiegel/Angebotsvergleich title, Supplier 1/2/3 columns, approval rows.
2. Quotation — Single vendor price offer: Vendor Name+Address, Date, Total Amount, Payment Terms, Delivery Time, Validity, Item Description. ONE vendor only.
3. RFQ — Request For Quotation (Motherson→vendors): Project Name, RFQ Number, Date, detailed specification table with "Supplier to specify" columns, Commercials section.
4. BER — Bid Exception Report: "BID EXCEPTION REPORT" header (required), Order Value, justification section, A-E checkbox options, three approval rows (Prepared/Purchasing/MD).
5. E-Auction — E-Auction results: Event ID, BID Id, Participant, Basic Price, Extended Price, eAuction/Reverse Auction terminology, bid timestamps.
6. Other — None of the above mandatory fields satisfied.

DECISION: Score each category by field coverage. Pick highest if ≥60% fields present. Otherwise: Other.
Confidence: 0.90+=all fields, 0.75-0.89=most fields, 0.60-0.74=enough, <0.60=Other.

OUTPUT (strict JSON, no markdown):
{"classification": "RFQ|Quotation|MPBC|BER|E-Auction|Other", "confidence": 0.0-1.0, "reason": "2-3 sentences", "key_signals": ["field1","field2","field3"]}"""

_CLASSIFY_USER_TEXT = """Classify this file. Return only JSON.

Filename: {filename}
Type: {file_type}

CONTENT:
{content}"""

_CLASSIFY_USER_IMAGE = """Classify this document image. Return only JSON.

Filename: {filename}
Type: image"""

_EXTRACT_SYSTEM = """You are a senior procurement analyst. Extract ALL line items from the quotation document.
Rules:
- Extract every item including continuation pages, addenda, optional items.
- Distinguish unit price vs total price. total = unit × quantity if not stated.
- Currency as ISO-4217 (INR, USD, EUR, GBP, ZAR, AED).
- Payment terms verbatim then normalized (e.g. "Net 30").
- If a field cannot be found, use null. Never hallucinate values.
- Translate all text to English.

Return ONLY a JSON array (no markdown fences, no commentary):
[{"item_name": "...", "quantity": number_or_null, "unit_price": number_or_null, "total_price": number_or_null, "currency": "...", "delivery_time": "...", "payment_terms": "...", "supplier_name": "..."}, ...]"""

_EXTRACT_USER = """Extract all line items from this quotation.

PR No: {pr_no}
File: {filename}

CONTENT:
{content}

Return the JSON array only."""

_BENCHMARK_SYSTEM = """You are a procurement price benchmark analyst.
Given a current item and historical similar items, recommend a fair benchmark unit price.
Return ONLY JSON: {"bp_unit_price": number_or_null, "summary": "2-3 sentence analysis covering price trend and recommendation"}"""

_BENCHMARK_USER = """Benchmark analysis:

CURRENT ITEM: {item_name}
  Quantity: {quantity}
  Currency: {currency}

HISTORICAL SIMILAR ITEMS ({n_hist} items):
{historical_json}

Return JSON only."""


class PipelineRunnerNode(Node):
    display_name = "Pipeline Runner"
    description = (
        "8-stage RAS attachment processing pipeline: Ingestion → "
        "Doc Download → Blob Upload → Classification → Metadata Extraction → "
        "Embeddings → Price Benchmark → Complete."
    )
    icon = "Workflow"
    name = "PipelineRunnerNode"

    inputs = [
        MessageInput(name="trigger", display_name="Trigger", advanced=True),
        DataInput(name="source_connection", display_name="Source DB (On-Prem)"),
        DataInput(name="target_connection", display_name="Target DB (Azure SQL)"),
        HandleInput(
            name="classify_llm",
            display_name="Classification LLM (GPT-4o-mini)",
            input_types=["LanguageModel"],
        ),
        HandleInput(
            name="extract_llm",
            display_name="Extraction LLM (GPT-4o)",
            input_types=["LanguageModel"],
        ),
        MessageTextInput(
            name="blob_connection_str",
            display_name="Azure Blob Connection String",
            value="",
            password=True,
        ),
        MessageTextInput(
            name="blob_container",
            display_name="Blob Container Name",
            value="ras-attachments",
        ),
        MessageTextInput(
            name="pinecone_api_key",
            display_name="Pinecone API Key (optional — enables embeddings+benchmark)",
            value="",
            password=True,
            advanced=True,
        ),
        MessageTextInput(
            name="pinecone_index",
            display_name="Pinecone Index Name",
            value="ras-quotations",
            advanced=True,
        ),
        MessageTextInput(
            name="embed_model_name",
            display_name="Embedding Model Name (Azure OpenAI deployment)",
            value="text-embedding-3-large",
            advanced=True,
        ),
        MessageTextInput(
            name="embed_api_key",
            display_name="Embedding API Key",
            value="",
            password=True,
            advanced=True,
        ),
        MessageTextInput(
            name="embed_endpoint",
            display_name="Embedding Azure OpenAI Endpoint",
            value="",
            advanced=True,
        ),
        MessageTextInput(
            name="pr_no_filter",
            display_name="Specific PR No (blank = all pending)",
            value="",
        ),
        IntInput(name="limit", display_name="Max PRs to Process", value=10),
        MessageTextInput(
            name="work_dir",
            display_name="Local Work Directory",
            value="/tmp/pipeline_work",
            advanced=True,
        ),
        MessageTextInput(
            name="az_schema",
            display_name="Azure SQL Schema",
            value="ras_procurement",
            advanced=True,
        ),
        MessageTextInput(
            name="src_schema",
            display_name="On-Prem Attachment Schema",
            value="ras_attachments",
            advanced=True,
        ),
        IntInput(
            name="classify_workers",
            display_name="Classification Parallel Workers",
            value=4,
            advanced=True,
        ),
    ]

    outputs = [
        Output(display_name="Pipeline Result", name="result", method="run_pipeline")
    ]

    # ─────────────────────────────────────────────────────────────────────────
    # Connection helpers
    # ─────────────────────────────────────────────────────────────────────────

    def _conn_str(self, conn_data: Data) -> str:
        d = conn_data.data or {}
        driver = d.get("driver", "ODBC Driver 18 for SQL Server")
        server = d.get("host", "")
        port = d.get("port", 1433)
        db = d.get("database_name", d.get("database", ""))
        user = d.get("username", "")
        pwd = d.get("password", "")
        return (
            f"DRIVER={{{driver}}};SERVER={server},{port};"
            f"DATABASE={db};UID={user};PWD={pwd};"
            f"TrustServerCertificate=yes;Connection Timeout=30;"
        )

    def _connect(self, conn_str: str) -> pyodbc.Connection:
        for attempt in range(4):
            try:
                return pyodbc.connect(conn_str, timeout=30)
            except Exception as exc:
                if attempt == 3:
                    raise
                time.sleep(2 ** attempt)

    # ─────────────────────────────────────────────────────────────────────────
    # Text / image extraction
    # ─────────────────────────────────────────────────────────────────────────

    def _extract_file(self, file_path: Path) -> tuple[str, Optional[str]]:
        """Returns (text_content, base64_image_or_None)."""
        ext = file_path.suffix.lower()
        try:
            file_bytes = file_path.read_bytes()
        except Exception:
            return "", None

        if ext == ".pdf":
            try:
                import fitz  # pymupdf
                doc = fitz.open(stream=file_bytes, filetype="pdf")
                text = "\n".join(page.get_text() for page in doc)
                if len(text.strip()) < 50 and len(doc) > 0:
                    pix = doc[0].get_pixmap(dpi=150)
                    b64 = base64.b64encode(pix.tobytes("png")).decode()
                    return text, b64
                return text[:15000], None
            except Exception:
                return "", None

        if ext in (".docx", ".doc"):
            try:
                import docx
                doc = docx.Document(io.BytesIO(file_bytes))
                return "\n".join(p.text for p in doc.paragraphs)[:15000], None
            except Exception:
                return "", None

        if ext in (".xlsx", ".xls"):
            try:
                import openpyxl
                wb = openpyxl.load_workbook(
                    io.BytesIO(file_bytes), read_only=True, data_only=True
                )
                rows = []
                for sheet in wb.worksheets:
                    rows.append(f"--- Sheet: {sheet.title} ---")
                    for row in sheet.iter_rows(max_row=500):
                        vals = [str(c.value) for c in row if c.value is not None]
                        if vals:
                            rows.append("\t".join(vals))
                return "\n".join(rows)[:15000], None
            except Exception:
                return "", None

        if ext == ".csv":
            return file_bytes.decode("utf-8", errors="replace")[:10000], None

        if ext in (".txt", ".html", ".htm"):
            return file_bytes.decode("utf-8", errors="replace")[:10000], None

        if ext in (".png", ".jpg", ".jpeg", ".tiff", ".tif", ".bmp"):
            b64 = base64.b64encode(file_bytes).decode()
            return "", b64

        return "", None

    # ─────────────────────────────────────────────────────────────────────────
    # Classification
    # ─────────────────────────────────────────────────────────────────────────

    _CLASSIFY_MAP = {"E-Auction": "E-Auction Results", "Other": "Others"}
    _IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".tiff", ".tif", ".bmp"}
    _SUPPORTED_EXTS = {
        ".pdf", ".docx", ".doc", ".xlsx", ".xls", ".csv",
        ".txt", ".html", ".htm",
        ".png", ".jpg", ".jpeg", ".tiff", ".tif", ".bmp",
    }

    def _is_trivial_image(self, file_path: Path) -> bool:
        if file_path.suffix.lower() not in self._IMAGE_EXTS:
            return False
        try:
            from PIL import Image
            with Image.open(file_path) as img:
                w, h = img.size
                if max(w, h) < 200 or min(w, h) < 100:
                    return True
                if max(w, h) / max(min(w, h), 1) > 5:
                    return True
                gray = img.convert("L")
                hist = gray.histogram()
                total = sum(hist) or 1
                if sum(hist[230:]) / total > 0.97:
                    return True
        except Exception:
            return True
        return False

    def _classify_file(self, file_path: Path) -> tuple[str, float]:
        """Returns (doc_type, confidence). Falls back to Others on any error."""
        if file_path.suffix.lower() not in self._SUPPORTED_EXTS:
            return "Others", 0.0
        if self._is_trivial_image(file_path):
            return "Others", 0.0

        text, b64 = self._extract_file(file_path)

        try:
            if b64 and not text.strip():
                user = _CLASSIFY_USER_IMAGE.format(filename=file_path.name)
                messages = [
                    SystemMessage(content=_CLASSIFY_SYSTEM),
                    HumanMessage(content=[
                        {"type": "text", "text": user},
                        {
                            "type": "image_url",
                            "image_url": {
                                "url": f"data:image/png;base64,{b64}",
                                "detail": "low",
                            },
                        },
                    ]),
                ]
            else:
                content = (text or "")[:8000]
                user = _CLASSIFY_USER_TEXT.format(
                    filename=file_path.name,
                    file_type=file_path.suffix,
                    content=content,
                )
                messages = [SystemMessage(content=_CLASSIFY_SYSTEM), HumanMessage(content=user)]

            resp = self.classify_llm.invoke(messages)
            raw = resp.content.strip()
            if raw.startswith("```"):
                parts = raw.split("```")
                raw = parts[1].lstrip("json").strip() if len(parts) > 1 else raw
            result = json.loads(raw)
            doc_type = self._CLASSIFY_MAP.get(
                result.get("classification", "Other"),
                result.get("classification", "Other"),
            )
            conf = float(result.get("confidence", 0.5))
            return doc_type, conf

        except Exception as exc:
            self.log(f"  Classification error for {file_path.name}: {exc}")
            return "Others", 0.0

    # ─────────────────────────────────────────────────────────────────────────
    # Metadata extraction
    # ─────────────────────────────────────────────────────────────────────────

    def _extract_items(self, file_path: Path, pr_no: str) -> list[dict]:
        text, _ = self._extract_file(file_path)
        content = (text or "")[:12000]
        if not content:
            return []

        user = _EXTRACT_USER.format(
            pr_no=pr_no, filename=file_path.name, content=content
        )
        try:
            resp = self.extract_llm.invoke(
                [SystemMessage(content=_EXTRACT_SYSTEM), HumanMessage(content=user)]
            )
            raw = resp.content.strip()
            if raw.startswith("```"):
                parts = raw.split("```")
                raw = parts[1].lstrip("json").strip() if len(parts) > 1 else raw
            items = json.loads(raw)
            return items if isinstance(items, list) else []
        except Exception as exc:
            self.log(f"  Extraction error for {file_path.name}: {exc}")
            return []

    # ─────────────────────────────────────────────────────────────────────────
    # Blob upload
    # ─────────────────────────────────────────────────────────────────────────

    def _upload_to_blob(self, work_pr_dir: Path) -> int:
        if not self.blob_connection_str.strip():
            return 0
        try:
            from azure.storage.blob import BlobServiceClient
            client = BlobServiceClient.from_connection_string(self.blob_connection_str)
            container = client.get_container_client(self.blob_container)
            count = 0
            base = work_pr_dir.parent.parent  # goes up to work_dir level
            for f in work_pr_dir.rglob("*"):
                if f.is_file():
                    blob_name = str(f.relative_to(base)).replace("\\", "/")
                    with open(f, "rb") as fh:
                        container.upload_blob(blob_name, fh, overwrite=True)
                    count += 1
            return count
        except Exception as exc:
            self.log(f"  Blob upload error: {exc}")
            return 0

    # ─────────────────────────────────────────────────────────────────────────
    # Embeddings (optional — requires Pinecone config)
    # ─────────────────────────────────────────────────────────────────────────

    def _run_embeddings(self, tgt_conn: pyodbc.Connection, pr_no: str) -> int:
        if not self.pinecone_api_key.strip() or not self.pinecone_index.strip():
            return 0
        if not self.embed_api_key.strip() or not self.embed_endpoint.strip():
            self.log("  Embedding API key / endpoint not set — skipping Pinecone upsert")
            return 0
        try:
            from openai import AzureOpenAI
            from pinecone import Pinecone

            schema = self.az_schema
            cur = tgt_conn.cursor()
            cur.execute(
                f"SELECT extracted_item_uuid_pk, item_name, purchase_dtl_id "
                f"FROM [{schema}].[quotation_extracted_items] "
                f"WHERE purchase_req_no = ? AND is_selected_quote = 1",
                pr_no,
            )
            rows = cur.fetchall()
            if not rows:
                return 0

            oai = AzureOpenAI(
                api_key=self.embed_api_key,
                azure_endpoint=self.embed_endpoint,
                api_version="2024-02-01",
            )
            pc = Pinecone(api_key=self.pinecone_api_key)
            index = pc.Index(self.pinecone_index)

            vectors = []
            for item_uuid, item_name, dtl_id in rows:
                if not item_name:
                    continue
                resp = oai.embeddings.create(
                    model=self.embed_model_name, input=item_name
                )
                vec = resp.data[0].embedding
                vectors.append({
                    "id": f"dtl_{dtl_id}",
                    "values": vec,
                    "metadata": {
                        "purchase_req_no": pr_no,
                        "purchase_dtl_id": int(dtl_id),
                        "item_name": item_name or "",
                    },
                })

            if vectors:
                # Upsert in batches of 100
                for i in range(0, len(vectors), 100):
                    index.upsert(vectors=vectors[i : i + 100])

            return len(vectors)

        except Exception as exc:
            self.log(f"  Embeddings error: {exc}")
            return 0

    # ─────────────────────────────────────────────────────────────────────────
    # Price benchmark (optional — requires Pinecone)
    # ─────────────────────────────────────────────────────────────────────────

    def _run_benchmark(self, tgt_conn: pyodbc.Connection, pr_no: str) -> int:
        if not self.pinecone_api_key.strip():
            return 0
        if not self.embed_api_key.strip() or not self.embed_endpoint.strip():
            return 0
        try:
            from openai import AzureOpenAI
            from pinecone import Pinecone

            schema = self.az_schema
            cur = tgt_conn.cursor()
            cur.execute(
                f"SELECT extracted_item_uuid_pk, item_name, quantity, currency, purchase_dtl_id "
                f"FROM [{schema}].[quotation_extracted_items] "
                f"WHERE purchase_req_no = ? AND is_selected_quote = 1",
                pr_no,
            )
            rows = cur.fetchall()
            if not rows:
                return 0

            oai = AzureOpenAI(
                api_key=self.embed_api_key,
                azure_endpoint=self.embed_endpoint,
                api_version="2024-02-01",
            )
            pc = Pinecone(api_key=self.pinecone_api_key)
            index = pc.Index(self.pinecone_index)

            written = 0
            for item_uuid, item_name, quantity, currency, dtl_id in rows:
                if not item_name:
                    continue

                # Embed current item
                emb_resp = oai.embeddings.create(
                    model=self.embed_model_name, input=item_name
                )
                vec = emb_resp.data[0].embedding

                # Query Pinecone for similar historical items (exclude current PR)
                qr = index.query(
                    vector=vec,
                    top_k=10,
                    filter={"purchase_req_no": {"$ne": pr_no}},
                    include_metadata=True,
                )
                historical = []
                for match in (qr.matches or []):
                    meta = match.metadata or {}
                    historical.append({
                        "item_name": meta.get("item_name", ""),
                        "purchase_req_no": meta.get("purchase_req_no", ""),
                        "score": round(match.score, 3),
                    })

                # LLM benchmark
                user = _BENCHMARK_USER.format(
                    item_name=item_name,
                    quantity=quantity,
                    currency=currency or "USD",
                    n_hist=len(historical),
                    historical_json=json.dumps(historical[:5], indent=2),
                )
                try:
                    resp = self.extract_llm.invoke(
                        [SystemMessage(content=_BENCHMARK_SYSTEM), HumanMessage(content=user)]
                    )
                    raw = resp.content.strip()
                    if raw.startswith("```"):
                        parts = raw.split("```")
                        raw = parts[1].lstrip("json").strip() if len(parts) > 1 else raw
                    bench = json.loads(raw)
                    bp_unit = bench.get("bp_unit_price")
                    summary = bench.get("summary", "")
                except Exception:
                    bp_unit, summary = None, ""

                cur2 = tgt_conn.cursor()
                cur2.execute(
                    f"""MERGE [{schema}].[benchmark_result] AS tgt
                        USING (SELECT ? AS purchase_dtl_id) AS src
                        ON tgt.purchase_dtl_id = src.purchase_dtl_id
                        WHEN MATCHED THEN UPDATE SET
                            bp_unit_price = ?, summary = ?, updated_at = SYSUTCDATETIME()
                        WHEN NOT MATCHED THEN INSERT
                            (extracted_item_uuid_fk, purchase_dtl_id, bp_unit_price, summary, created_at, updated_at)
                            VALUES (?, ?, ?, ?, SYSUTCDATETIME(), SYSUTCDATETIME());""",
                    dtl_id, bp_unit, summary, str(item_uuid), dtl_id, bp_unit, summary,
                )
                tgt_conn.commit()
                written += 1

            return written

        except Exception as exc:
            self.log(f"  Benchmark error: {exc}")
            return 0

    # ─────────────────────────────────────────────────────────────────────────
    # DB helpers
    # ─────────────────────────────────────────────────────────────────────────

    def _upsert_tracker(
        self, tgt_conn: pyodbc.Connection, pr_no: str, stage_id: int
    ) -> None:
        schema = self.az_schema
        cur = tgt_conn.cursor()
        cur.execute(
            f"""MERGE [{schema}].[ras_tracker] AS tgt
                USING (SELECT ? AS purchase_req_no) AS src
                ON tgt.purchase_req_no = src.purchase_req_no
                WHEN MATCHED THEN UPDATE SET
                    current_stage_fk = ?, updated_at = SYSUTCDATETIME()
                WHEN NOT MATCHED THEN INSERT
                    (purchase_req_no, current_stage_fk, created_at, updated_at)
                    VALUES (?, ?, SYSUTCDATETIME(), SYSUTCDATETIME());""",
            pr_no, stage_id, pr_no, stage_id,
        )
        tgt_conn.commit()

    def _fetch_pending_prs(self, tgt_cs: str) -> list[str]:
        schema = self.az_schema
        conn = self._connect(tgt_cs)
        cur = conn.cursor()
        try:
            cur.execute(
                f"""SELECT TOP (?) m.PURCHASE_REQ_NO
                    FROM [{schema}].[purchase_req_mst] m
                    LEFT JOIN [{schema}].[ras_tracker] t
                        ON t.purchase_req_no = m.PURCHASE_REQ_NO
                    WHERE t.purchase_req_no IS NULL
                       OR t.current_stage_fk < 8
                    ORDER BY m.C_DATETIME""",
                self.limit,
            )
            return [r[0] for r in cur.fetchall()]
        finally:
            conn.close()

    def _download_attachments(
        self, src_cs: str, pr_no: str, work_pr_dir: Path
    ) -> int:
        schema = self.src_schema
        conn = self._connect(src_cs)
        cur = conn.cursor()
        try:
            cur.execute(
                f"""SELECT d.ATTACHMENT_ID, d.ATTACHMENT_NAME, d.ATTACHMENT_DATA
                    FROM [{schema}].[RAS_ATTACHMENT_DOCUMENTS] d
                    JOIN [{schema}].[RAS_ATTACHMENTS] a
                        ON a.RAS_ID = d.ATTACHMENT_ID
                    WHERE a.PURCHASE_REQ_NO = ?""",
                pr_no,
            )
            rows = cur.fetchall()
        finally:
            conn.close()

        count = 0
        for att_id, att_name, att_data in rows:
            if att_data is None:
                continue
            att_dir = work_pr_dir / str(att_id)
            att_dir.mkdir(parents=True, exist_ok=True)
            safe_name = Path(att_name).name if att_name else f"{att_id}.bin"
            (att_dir / safe_name).write_bytes(att_data)
            count += 1
        return count

    def _get_or_create_tracker_uuid(
        self, tgt_conn: pyodbc.Connection, pr_no: str
    ) -> str:
        schema = self.az_schema
        cur = tgt_conn.cursor()
        cur.execute(
            f"SELECT rass_uuid_pk FROM [{schema}].[ras_tracker] WHERE purchase_req_no = ?",
            pr_no,
        )
        row = cur.fetchone()
        return str(row[0]) if row and row[0] else str(uuid.uuid4())

    def _upsert_parent_classification(
        self,
        tgt_conn: pyodbc.Connection,
        pr_no: str,
        rass_uuid: str,
        att_id: str,
        blob_path: str,
        doc_type: str,
        conf: float,
    ) -> None:
        schema = self.az_schema
        cur = tgt_conn.cursor()
        cur.execute(
            f"""MERGE [{schema}].[attachment_classification] AS tgt
                USING (SELECT ? AS attachment_id) AS src
                ON tgt.attachment_id = src.attachment_id
                WHEN MATCHED THEN UPDATE SET
                    doc_type = ?, classification_conf = ?,
                    blob_file_path = ?, updated_at = SYSUTCDATETIME()
                WHEN NOT MATCHED THEN INSERT
                    (ras_uuid_fk, purchase_req_no, attachment_id,
                     blob_file_path, doc_type, classification_conf,
                     created_at, updated_at)
                    VALUES (?, ?, ?, ?, ?, ?,
                            SYSUTCDATETIME(), SYSUTCDATETIME());""",
            att_id, doc_type, conf, blob_path,
            rass_uuid, pr_no, att_id, blob_path, doc_type, conf,
        )
        tgt_conn.commit()

    def _write_extracted_items(
        self,
        tgt_conn: pyodbc.Connection,
        pr_no: str,
        items: list[dict],
        blob_path: str,
    ) -> None:
        schema = self.az_schema
        cur = tgt_conn.cursor()
        for item in items:
            cur.execute(
                f"""INSERT INTO [{schema}].[quotation_extracted_items]
                    (purchase_req_no, blob_file_path, item_name, quantity,
                     unit_price, total_price, currency, delivery_time,
                     payment_terms, supplier_name, created_at)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, SYSUTCDATETIME())""",
                pr_no,
                blob_path,
                item.get("item_name"),
                item.get("quantity"),
                item.get("unit_price"),
                item.get("total_price"),
                item.get("currency"),
                str(item.get("delivery_time") or "")[:200],
                str(item.get("payment_terms") or "")[:200],
                item.get("supplier_name"),
            )
        tgt_conn.commit()

    # ─────────────────────────────────────────────────────────────────────────
    # Per-PR pipeline runner
    # ─────────────────────────────────────────────────────────────────────────

    def _process_pr(self, pr_no: str, src_cs: str, tgt_cs: str) -> dict:
        result = {
            "pr_no": pr_no,
            "status": "FAILED",
            "stages": {},
            "error": "",
        }

        safe_pr = pr_no.replace("/", "_")
        work_pr_dir = Path(self.work_dir) / "procurement" / safe_pr
        if work_pr_dir.exists():
            shutil.rmtree(work_pr_dir)
        work_pr_dir.mkdir(parents=True, exist_ok=True)

        tgt_conn = self._connect(tgt_cs)
        try:
            # ── Stage 1: Ingestion ──────────────────────────────────────────
            self.log(f"[{pr_no}] Stage 1: Ingestion")
            self._upsert_tracker(tgt_conn, pr_no, 1)
            result["stages"]["INGESTION"] = "OK"

            # ── Stage 2: Download attachments ──────────────────────────────
            self.log(f"[{pr_no}] Stage 2: Doc download + extraction")
            downloaded = self._download_attachments(src_cs, pr_no, work_pr_dir)
            self.log(f"[{pr_no}]   → {downloaded} file(s) saved")
            self._upsert_tracker(tgt_conn, pr_no, 2)
            result["stages"]["EMBED_DOC_EXTRACTION"] = f"OK ({downloaded} files)"

            # ── Stage 3: Blob upload ────────────────────────────────────────
            self.log(f"[{pr_no}] Stage 3: Blob upload")
            uploaded = self._upload_to_blob(work_pr_dir)
            self.log(f"[{pr_no}]   → {uploaded} file(s) uploaded")
            self._upsert_tracker(tgt_conn, pr_no, 3)
            result["stages"]["BLOB_UPLOAD"] = f"OK ({uploaded} files)"

            # ── Stage 4: Classification ─────────────────────────────────────
            self.log(f"[{pr_no}] Stage 4: Classification")
            rass_uuid = self._get_or_create_tracker_uuid(tgt_conn, pr_no)
            quotation_files: list[Path] = []
            classified = 0

            # Parallel classification using ThreadPoolExecutor
            from concurrent.futures import ThreadPoolExecutor, as_completed
            file_tasks = []
            for att_dir in sorted(work_pr_dir.iterdir()):
                if not att_dir.is_dir():
                    continue
                for f in sorted(att_dir.iterdir()):
                    if f.is_file():
                        file_tasks.append((f, att_dir.name))

            workers = max(1, min(self.classify_workers, len(file_tasks)))
            with ThreadPoolExecutor(max_workers=workers) as pool:
                futures = {
                    pool.submit(self._classify_file, f): (f, att_id)
                    for f, att_id in file_tasks
                }
                for fut in as_completed(futures):
                    f, att_id = futures[fut]
                    try:
                        doc_type, conf = fut.result()
                    except Exception as exc:
                        self.log(f"  classify error {f.name}: {exc}")
                        doc_type, conf = "Others", 0.0
                    blob_path = f"procurement/{safe_pr}/{att_id}/{f.name}"
                    self._upsert_parent_classification(
                        tgt_conn, pr_no, rass_uuid, att_id, blob_path, doc_type, conf
                    )
                    if doc_type == "Quotation":
                        quotation_files.append(f)
                    self.log(f"[{pr_no}]   {f.name} → {doc_type} ({conf:.2f})")
                    classified += 1

            self._upsert_tracker(tgt_conn, pr_no, 4)
            result["stages"]["CLASSIFICATION"] = (
                f"OK ({classified} files, {len(quotation_files)} quotations)"
            )

            # ── Stage 5: Metadata extraction ────────────────────────────────
            self.log(f"[{pr_no}] Stage 5: Metadata extraction")
            total_items = 0
            for qf in quotation_files:
                att_id = qf.parent.name
                blob_path = f"procurement/{safe_pr}/{att_id}/{qf.name}"
                items = self._extract_items(qf, pr_no)
                if items:
                    self._write_extracted_items(tgt_conn, pr_no, items, blob_path)
                    total_items += len(items)
                    self.log(f"[{pr_no}]   {qf.name} → {len(items)} item(s)")

            if not quotation_files:
                self.log(f"[{pr_no}]   No quotation files — metadata extraction skipped")
            self._upsert_tracker(tgt_conn, pr_no, 5)
            result["stages"]["METADATA_EXTRACTION"] = f"OK ({total_items} items)"

            # ── Stage 6: Embeddings ──────────────────────────────────────────
            self.log(f"[{pr_no}] Stage 6: Embeddings")
            emb_count = self._run_embeddings(tgt_conn, pr_no)
            emb_status = (
                f"OK ({emb_count} vectors)" if emb_count
                else "SKIPPED (no Pinecone/embedding config)"
            )
            self._upsert_tracker(tgt_conn, pr_no, 6)
            result["stages"]["EMBEDDINGS"] = emb_status

            # ── Stage 7: Price benchmark ─────────────────────────────────────
            self.log(f"[{pr_no}] Stage 7: Price benchmark")
            bench_count = self._run_benchmark(tgt_conn, pr_no)
            bench_status = (
                f"OK ({bench_count} rows)" if bench_count
                else "SKIPPED (no Pinecone/embedding config)"
            )
            self._upsert_tracker(tgt_conn, pr_no, 7)
            result["stages"]["PRICE_BENCHMARK"] = bench_status

            # ── Stage 8: Complete ────────────────────────────────────────────
            self.log(f"[{pr_no}] Stage 8: Complete")
            self._upsert_tracker(tgt_conn, pr_no, 8)
            result["stages"]["COMPLETE"] = "OK"
            result["status"] = "SUCCESS"

        except Exception as exc:
            result["error"] = str(exc)
            self.log(f"[{pr_no}] FAILED at stage: {exc}")

        finally:
            tgt_conn.close()
            # Clean up local work folder regardless of success/failure
            try:
                if work_pr_dir.exists():
                    shutil.rmtree(work_pr_dir)
            except Exception:
                pass

        return result

    # ─────────────────────────────────────────────────────────────────────────
    # Entry point
    # ─────────────────────────────────────────────────────────────────────────

    def run_pipeline(self) -> Message:
        src_cs = self._conn_str(self.source_connection)
        tgt_cs = self._conn_str(self.target_connection)

        # Resolve which PRs to run
        if self.pr_no_filter.strip():
            pr_nos = [self.pr_no_filter.strip()]
            self.log(f"Single-PR mode: {pr_nos[0]}")
        else:
            self.log("Fetching pending PRs from DB...")
            pr_nos = self._fetch_pending_prs(tgt_cs)
            self.log(f"Found {len(pr_nos)} pending PR(s)")

        if not pr_nos:
            return Message(text="Pipeline complete — no pending PRs to process.")

        results = []
        for idx, pr_no in enumerate(pr_nos, 1):
            self.log(f"\n--- [{idx}/{len(pr_nos)}] Processing: {pr_no} ---")
            r = self._process_pr(pr_no, src_cs, tgt_cs)
            results.append(r)

        ok = [r for r in results if r["status"] == "SUCCESS"]
        fail = [r for r in results if r["status"] == "FAILED"]

        lines = [
            f"Pipeline Run Complete",
            f"  Success: {len(ok)}/{len(results)} PRs",
            "",
        ]
        for r in ok:
            stage_summary = " → ".join(
                f"{k}:{v.split('(')[0].strip()}" for k, v in r["stages"].items()
            )
            lines.append(f"  OK  [{r['pr_no']}]  {stage_summary}")
        for r in fail:
            lines.append(f"  ERR [{r['pr_no']}]  {r['error']}")

        return Message(text="\n".join(lines))
